import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from datetime import datetime
from typing import Dict
import os
from config import config

class ExcelService:
    def __init__(self):
        self.excel_path = config.EXCEL_FILE_PATH
        self._ensure_workbook_exists()

    def _ensure_workbook_exists(self):
        """Create Excel workbook if it doesn't exist"""
        if not os.path.exists(self.excel_path):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Job Applications"

            # Define headers
            headers = [
                "Date Added",
                "Company",
                "Job Title",
                "Role Type",
                "Compensation",
                "Job Description",
                "Application Link",
                "Application Method",
                "Email Subject",
                "Resume Link",
                "Status"
            ]

            # Write headers
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.value = header
                cell.font = Font(bold=True, size=12)
                cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                cell.font = Font(bold=True, size=12, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Set column widths
            column_widths = {
                "A": 15,  # Date
                "B": 20,  # Company
                "C": 25,  # Job Title
                "D": 12,  # Role Type
                "E": 15,  # Compensation
                "F": 60,  # Job Description
                "G": 40,  # Application Link
                "H": 15,  # Application Method
                "I": 30,  # Email Subject
                "J": 50,  # Resume Link
                "K": 12,  # Status
            }

            for col, width in column_widths.items():
                ws.column_dimensions[col].width = width

            # Enable text wrapping for description column
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=6, max_col=6):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

            wb.save(self.excel_path)

    def add_job_application(self, job_data: Dict, resume_link: str = "") -> int:
        """
        Add a new job application to the Excel sheet
        Returns the row number where data was added
        """
        try:
            wb = openpyxl.load_workbook(self.excel_path)
            ws = wb.active

            # Find next empty row
            next_row = ws.max_row + 1

            # Prepare data
            row_data = [
                datetime.now().strftime("%Y-%m-%d"),
                job_data.get("company", ""),
                job_data.get("job_title", ""),
                job_data.get("role_type", ""),
                job_data.get("compensation", ""),
                job_data.get("job_description", ""),
                job_data.get("application_link", ""),
                job_data.get("application_method", ""),
                job_data.get("email_subject", ""),
                resume_link,
                "New"
            ]

            # Write data
            for col_num, value in enumerate(row_data, 1):
                cell = ws.cell(row=next_row, column=col_num)
                cell.value = value

                # Wrap text for description
                if col_num == 6:  # Job Description column
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

                # Make resume link clickable if present
                if col_num == 10 and value:  # Resume Link column
                    cell.hyperlink = value
                    cell.font = Font(color="0000FF", underline="single")

                # Make application link clickable
                if col_num == 7 and value and value.startswith("http"):  # Application Link
                    cell.hyperlink = value
                    cell.font = Font(color="0000FF", underline="single")

            wb.save(self.excel_path)
            return next_row

        except Exception as e:
            raise Exception(f"Error adding to Excel: {str(e)}")

    def update_resume_link(self, row_number: int, resume_link: str):
        """Update the resume link for a specific row"""
        try:
            wb = openpyxl.load_workbook(self.excel_path)
            ws = wb.active

            # Update resume link (column J = 10)
            cell = ws.cell(row=row_number, column=10)
            cell.value = resume_link
            cell.hyperlink = resume_link
            cell.font = Font(color="0000FF", underline="single")

            # Update status
            status_cell = ws.cell(row=row_number, column=11)
            status_cell.value = "Resume Generated"

            wb.save(self.excel_path)

        except Exception as e:
            raise Exception(f"Error updating resume link: {str(e)}")

    def get_excel_path(self) -> str:
        """Return the path to the Excel file"""
        return os.path.abspath(self.excel_path)
